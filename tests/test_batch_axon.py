"""Batch_Axon workbooks: formulas, missing values, which animals get a workbook."""
from pathlib import Path

import openpyxl
import pytest

import batch_axon.compile as compile_mod
from batch_axon.compile import compile_study

ROWS = [
    {"name": "img_1", "csa_um2": 1000.0, "total_axons": 50, "gratio": 0.62, "axon_diam_um": 3.0, "axon_density_per_um2": 0.05},
    {"name": "img_2", "csa_um2": None,   "total_axons": 0,  "gratio": None, "axon_diam_um": None, "axon_density_per_um2": None},
]
AGG = {"fourx_csa_um2": 50000.0, "total_images": 2, "total_axons": 50, "mean_gratio": 0.62,
       "mean_axon_diam_um": 3.0, "total_sample_csa_um2": None, "estimated_total_axons": None}


def _study(root: Path, layout: dict) -> dict:
    """layout {animal: {nerve: has_morphometrics}} -> the dict resolve_scan() returns."""
    study = {}
    for animal, nerves in layout.items():
        study[animal] = {}
        for nerve, has_morph in nerves.items():
            tiff = root / animal / nerve / "40X_TIFF"
            tiff.mkdir(parents=True)
            study[animal][nerve] = {"tiff_dir": tiff, "mag": "40X",
                                    "morphometrics_dir": (tiff.parent / "Morphometrics") if has_morph else None}
    return study


@pytest.fixture
def study(tmp_path, monkeypatch):
    monkeypatch.setattr(compile_mod, "get_nerve_data", lambda nerve_dir, mag, log: (ROWS, AGG, None))
    root = tmp_path / "Study"
    return root, _study(root, {"A1": {"N1": True, "N2": False}, "A2": {"N1": False}})


def test_workbooks_written_and_counts(study, log):
    root, s = study
    r = compile_study(s, root, "40X", {}, log)
    assert (root / "A1" / "A1_Data.xlsx").exists()
    assert not (root / "A2" / "A2_Data.xlsx").exists()        # no nerve data: no workbook kept
    assert r.combined_path == root / "Study_Data.xlsx"
    assert (r.animals_compiled, r.animals_skipped, r.nerves_processed, r.nerves_failed, r.nerves_incomplete) == (1, 1, 1, 2, 1)
    assert openpyxl.load_workbook(r.combined_path).sheetnames == ["Provenance", "A1", "A2"]


def test_nerve_block_formulas_and_missing(study, log):
    root, s = study
    compile_study(s, root, "40X", {}, log)
    ws = openpyxl.load_workbook(root / "A1" / "A1_Data.xlsx").active
    assert ws["A2"].value == "A1 - N1" and ws["C2"].value == 50000.0          # 4X whole-nerve CSA
    assert [ws.cell(3, c).value for c in (2, 3, 4)] == ["img_1", 1000.0, 50]
    assert [ws.cell(4, c).value for c in (2, 3, 4, 5, 6)] == ["img_2", "missing", 0, "missing", "missing"]
    assert ws["H3"].value.startswith("=IF(OR(NOT(ISNUMBER(C3)),NOT(ISNUMBER(D3)))")
    assert ws["B5"].value == "Totals (n=2)"
    assert ws["C5"].value == '=IF(COUNT(C3:C4)=2,SUM(C3:C4),"missing")'    # incomplete CSA -> missing
    assert ws["G5"].value == '=IF(OR(NOT(ISNUMBER(C5)),NOT(ISNUMBER(C2))),"missing",IFERROR(D5/C5*C2,"missing"))'
    ws_study = openpyxl.load_workbook(root / "Study_Data.xlsx")["A1"]
    assert [[c.value for c in row] for row in ws_study.iter_rows()] == [[c.value for c in row] for row in ws.iter_rows()]


def test_compile_only_copies_values(study, log):
    root, s = study
    compile_study(s, root, "40X", {}, log)
    r = compile_mod.compile_from_animal_workbooks(s, root, "40X", {}, log)
    assert (r.animals_compiled, r.animals_skipped) == (1, 1)
    ws = openpyxl.load_workbook(r.combined_path)["A1"]
    assert ws["B4"].value == "img_2" and ws["C4"].value == "missing"
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)


def _block(tmp_path, rows, fourx=50000.0):
    """Write one nerve block; return (formulas sheet, stored-values sheet)."""
    import xlsxwriter
    from batch_axon.workbook import add_formats, write_nerve_block
    p = tmp_path / "block.xlsx"
    wb = xlsxwriter.Workbook(str(p))
    write_nerve_block(wb.add_worksheet("A1"), 0, "N1", "A1", rows, {"fourx_csa_um2": fourx}, add_formats(wb))
    wb.close()
    return openpyxl.load_workbook(p).active, openpyxl.load_workbook(p, data_only=True).active


def test_totals_weighted_by_axon_count(tmp_path):
    rows = [{"name": "a", "csa_um2": 1000.0, "total_axons": 100, "gratio": 0.6, "axon_diam_um": 2.0},
            {"name": "b", "csa_um2": 3000.0, "total_axons": 300, "gratio": 0.7, "axon_diam_um": 4.0}]
    f, v = _block(tmp_path, rows)
    assert f["E5"].value.startswith("=IF(SUMPRODUCT(D3:D4,--ISNUMBER(E3:E4))>0")
    totals = [v.cell(5, c).value for c in range(3, 9)]
    assert totals == pytest.approx([4000.0, 400, 0.675, 3.5, 5000.0, 0.1])   # pooled, not (0.6+0.7)/2


def test_totals_stored_values_follow_formulas_when_csa_missing(tmp_path):
    # Stored values are what Excel computes: a missing CSA makes CSA, density
    # and the estimate "missing" (not a partial sum); means skip the image
    # with no g-ratio; a zero-axon image keeps its 0.
    f, v = _block(tmp_path, ROWS)
    assert [v.cell(4, c).value for c in (3, 4, 5, 6, 8)] == ["missing", 0, "missing", "missing", "missing"]
    assert [v.cell(5, c).value for c in range(3, 9)] == ["missing", 50, 0.62, 3.0, "missing", "missing"]
    assert v["H3"].value == pytest.approx(0.05)


def test_outside_tracing_columns(tmp_path):
    # Images 1 and 2 had a tracing applied (3 and 1 axons outside it); image 3 did not.
    rows = [{"name": "a", "csa_um2": 1000.0, "total_axons": 97, "gratio": 0.6, "axon_diam_um": 2.0, "axons_outside_tracing": 3},
            {"name": "b", "csa_um2": 1000.0, "total_axons": 99, "gratio": 0.6, "axon_diam_um": 2.0, "axons_outside_tracing": 1},
            {"name": "c", "csa_um2": 1000.0, "total_axons": 50, "gratio": 0.6, "axon_diam_um": 2.0}]
    f, v = _block(tmp_path, rows)
    assert [v.cell(2 + 1, c).value for c in (9, 10)] == [3, pytest.approx(0.03)]
    assert [v.cell(5, c).value for c in (9, 10)] == ["n/a", "n/a"]          # image c: no tracing applied
    assert [v.cell(6, c).value for c in (9, 10)] == [4, pytest.approx(4 / (4 + 97 + 99))]   # c is left out of the share
    assert f["I6"].value == '=IF(COUNT(I3:I5)=0,"n/a",SUM(I3:I5))'
    assert [c.value for c in f[1]][8:10] == ["Axons Outside Tracing", "Outside Tracing (%)"]


def test_outside_tracing_columns_without_any_tracing(tmp_path):
    f, v = _block(tmp_path, ROWS)
    assert [v.cell(3, c).value for c in (9, 10)] == ["n/a", "n/a"]
    assert [v.cell(5, c).value for c in (9, 10)] == ["n/a", "n/a"]