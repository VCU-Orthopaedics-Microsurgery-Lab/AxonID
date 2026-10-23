"""Whole-image exclusion: class EXCLUDE + a reason; nothing in the image is counted, and it is listed."""
import json

import openpyxl
import pytest
import tifffile
import xlsxwriter

from batch_axon.overlay.process_overlay_qupath import get_geojson_exclusion, get_geojson_region
from batch_axon.workbook import add_formats, write_nerve_block
from conftest import draw_fibers
from morphometrics.analyze_nerve import get_nerve_data
from morphometrics.distributions import bin_nerve_diameters, image_exclusion_reason
from morphometrics.morphometrics import get_morphometrics, save_morphometrics
from morphometrics.tracing import exclude_image
from utils.calibration import Calibration

CONFIG = """
[pixel_size_um.40X]
1440 = 0.2

[folders]
tiff_suffixes = ["_TIFF"]
csa_suffix_geojson = "_CSA.geojson"
"""
CENTRES = [(35, 35), (35, 85), (85, 35), (85, 85)]


def feature(cls, name=None, side=500):
    poly = [[[0, 0], [side, 0], [side, side], [0, side], [0, 0]]]
    props = {"classification": {"name": cls}}
    if name is not None:
        props["name"] = name
    return {"type": "Feature", "properties": props, "geometry": {"type": "Polygon", "coordinates": poly}}


def write_geojson(path, features):
    path.write_text(json.dumps({"type": "FeatureCollection", "features": features}))


class Log:
    def __init__(self):
        self.warnings, self.infos = [], []

    def warn(self, m): self.warnings.append(m)
    def info(self, m): self.infos.append(m)
    def error(self, m): self.warnings.append(m)
    def success(self, m): pass


# ── reading the tag ──────────────────────────────────────────────────────────

def test_reason_is_the_name_and_class_is_case_insensitive(tmp_path):
    p = tmp_path / "a.geojson"
    write_geojson(p, [feature("+"), feature("Exclude", "out of focus")])
    assert get_geojson_exclusion(str(p)) == "out of focus"


def test_tag_without_reason_is_refused(tmp_path, capsys):
    p = tmp_path / "a.geojson"
    write_geojson(p, [feature("+"), feature("EXCLUDE")])
    assert get_geojson_exclusion(str(p)) is None                 # NOT excluded
    assert "no reason" in capsys.readouterr().out
    write_geojson(p, [feature("EXCLUDE", "   ")])
    assert get_geojson_exclusion(str(p), quiet=True) is None


def test_description_is_a_fallback_and_reasons_join(tmp_path):
    p = tmp_path / "a.geojson"
    f2 = feature("EXCLUDE")
    f2["properties"]["description"] = "folded tissue"
    write_geojson(p, [feature("EXCLUDE", "blurry"), f2, feature("EXCLUDE", "blurry")])
    assert get_geojson_exclusion(str(p)) == "blurry; folded tissue"


def test_no_tag_no_exclusion(tmp_path):
    p = tmp_path / "a.geojson"
    write_geojson(p, [feature("+"), feature("-")])
    assert get_geojson_exclusion(str(p)) is None
    assert get_geojson_exclusion(str(tmp_path / "missing.geojson")) is None


def test_region_ignores_the_tag_silently(tmp_path, capsys):
    p = tmp_path / "a.geojson"
    write_geojson(p, [feature("+", side=100), feature("EXCLUDE", "blurry")])
    assert get_geojson_region(str(p)).area == pytest.approx(100 * 100)
    assert capsys.readouterr().out == ""                         # no "untagged feature" warning


# ── excluding an image ───────────────────────────────────────────────────────

def test_exclude_image_returns_everything_as_outside(tmp_path, temp_config, log):
    temp_config(CONFIG)
    seg = tmp_path / "s_segmented.tif"
    tifffile.imwrite(seg, draw_fibers((120, 120), CENTRES, 15, 9))
    cal = Calibration(0.2, 120, 120, 1.0, "provenance")
    df, px = get_morphometrics(str(seg), "40X", log, calibration=cal)
    kept, kept_px, outside, info = exclude_image(df, px, 0.5, "folded")
    assert kept.empty and kept_px.empty and len(outside) == len(df) == 4
    assert set(outside["excluded_because"]) == {"image excluded: folded"}
    assert outside["qupath_x"].iloc[0] == pytest.approx((df["y"].iloc[0] + 0.5) / 0.5, abs=0.1)
    assert info == {"n_all": 4, "n_outside": 4, "region_px": None, "excluded": "folded"}


# ── nerve level: not in any number, listed with the reason ───────────────────

def build_nerve(tmp_path, log, exclude_second=True):
    nerve = tmp_path / "N1"
    (nerve / "CSA").mkdir(parents=True)
    for name in ("img_1", "img_2"):
        seg = tmp_path / f"{name}_segmented.tif"
        tifffile.imwrite(seg, draw_fibers((120, 120), CENTRES, 15, 9))
        cal = Calibration(0.2, 1440, 1024, 1.0, "provenance")
        df, px = get_morphometrics(str(seg), "40X", log, calibration=cal)
        info = None
        if name == "img_2" and exclude_second:
            df, px, outside, info = exclude_image(df, px, 1.0, "out of focus")
        save_morphometrics(df, px, str(nerve / "Morphometrics"), name, calibration=cal, tracing=info,
                           outside_df=outside if info else None)
        write_geojson(nerve / "CSA" / f"{name}_CSA.geojson", [feature("+")])
    return nerve


def test_excluded_image_is_in_no_number(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve = build_nerve(tmp_path, log)
    rows, agg, _ = get_nerve_data(nerve, "40X", log)
    assert [r["name"] for r in rows] == ["img_1"]
    assert agg["total_images"] == 1 and agg["total_axons"] == 4
    assert agg["excluded_images"] == [{"name": "img_2", "reason": "out of focus"}]
    assert agg["total_sample_csa_um2"] is None or agg["total_sample_csa_um2"] == pytest.approx(10000.0)
    bins = bin_nerve_diameters(nerve / "Morphometrics", "N1", "40X", log)
    assert list(bins["per_image_df"]["Image"]) == ["img_1"]
    assert image_exclusion_reason(nerve / "Morphometrics" / "img_2_morphometrics.xlsx") == "out of focus"
    assert image_exclusion_reason(nerve / "Morphometrics" / "img_1_morphometrics.xlsx") is None


def test_summary_sheet_records_the_exclusion(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve = build_nerve(tmp_path, log)
    wb = openpyxl.load_workbook(nerve / "Morphometrics" / "img_2_morphometrics.xlsx")
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(max_col=2) if r[0].value}
    assert summary["Image excluded"] == "yes: out of focus"
    assert summary["CSA tracing applied"] == "no (image excluded)"
    assert wb["Outside tracing"].max_row == 1 + 4                 # the 4 axons are kept, listed


def test_tag_added_after_morphometrics_warns(tmp_path, temp_config):
    temp_config(CONFIG)
    log = Log()
    nerve = build_nerve(tmp_path, log, exclude_second=False)
    write_geojson(nerve / "CSA" / "img_2_CSA.geojson", [feature("+"), feature("EXCLUDE", "blurry")])
    rows, agg, _ = get_nerve_data(nerve, "40X", log)
    assert [r["name"] for r in rows] == ["img_1", "img_2"]         # still counted: not applied yet
    assert any("tagged EXCLUDE" in w and "img_2" in w for w in log.warnings)


def test_data_workbook_lists_excluded_images_outside_the_totals(tmp_path):
    rows = [{"name": "a", "csa_um2": 1000.0, "total_axons": 100, "gratio": 0.6, "axon_diam_um": 2.0}]
    agg = {"fourx_csa_um2": 50000.0, "excluded_images": [{"name": "b", "reason": "folded tissue"}]}
    p = tmp_path / "w.xlsx"
    book = xlsxwriter.Workbook(str(p))
    write_nerve_block(book.add_worksheet("A1"), 0, "N1", "A1", rows, agg, add_formats(book))
    book.close()
    ws = openpyxl.load_workbook(p).active
    assert ws["B4"].value == "Totals (n=1)"                        # the excluded image is not in n
    assert ws["B6"].value == "b" and ws["C6"].value == "EXCLUDED: folded tissue"
    assert "Excluded images (n=1)" in ws["B5"].value