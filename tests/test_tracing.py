"""Axons outside a CSA tracing are not measured, but they are counted and listed."""
import json

import openpyxl
import pytest
import tifffile

from batch_axon.overlay.process_overlay_qupath import get_geojson_region, points_in_region
from conftest import draw_fibers
from morphometrics.analyze_nerve import get_nerve_data
from morphometrics.distributions import bin_nerve_diameters
from morphometrics.morphometrics import get_morphometrics, save_morphometrics
from morphometrics.tracing import split_by_tracing, tracing_path
from utils.calibration import Calibration

# Image QuPath traced: 120 px wide at 0.2 µm/px (segmented at the same size).
CONFIG = """
[pixel_size_um.40X]
120 = 0.2

[folders]
tiff_suffixes = ["_TIFF"]
csa_suffix_geojson = "_CSA.geojson"
"""
# Same, but the traced image is 240 px wide (0.1 µm/px) and was segmented at 120 px.
CONFIG_DOWNSAMPLED = CONFIG.replace("120 = 0.2", "240 = 0.1")


def _geojson(path, features):
    path.write_text(json.dumps({"type": "FeatureCollection", "features": features}))


def _box(x0, y0, x1, y1, name):
    ring = [[x0, y0], [x1, y0], [x1, y1], [x0, y1], [x0, y0]]
    return {"type": "Feature", "properties": {"name": name},
            "geometry": {"type": "Polygon", "coordinates": [ring]}}


# Four fibers at (row, column): columns 35 (left) and 85 (right), rows 35 and 85.
CENTRES = [(35, 35), (35, 85), (85, 35), (85, 85)]


def _counts(info):
    return info["n_all"], info["n_outside"]


def _measure(tmp_path, log, features=None, apply=True, scale=1.0, name="img_1"):
    """Segmented image -> morphometrics workbook, as the morphometrics stage does it."""
    nerve = tmp_path / "N1"
    (nerve / "CSA").mkdir(parents=True, exist_ok=True)
    seg = tmp_path / f"{name}_segmented.tif"
    tifffile.imwrite(seg, draw_fibers((120, 120), CENTRES, 15, 9))
    orig = int(120 / scale)
    cal = Calibration(0.2, orig, orig, scale, "provenance")
    df, raw_df = get_morphometrics(str(seg), "40X", log, calibration=cal)
    info = outside = None
    if features is not None:
        gj = nerve / "CSA" / f"{name}_CSA.geojson"
        _geojson(gj, features)
        if apply:
            df, raw_df, outside, info = split_by_tracing(df, raw_df, gj, scale, log)
    save_morphometrics(df, raw_df, str(nerve / "Morphometrics"), name,
                       calibration=cal, tracing=info, outside_df=outside)
    return nerve, df, outside, info


def test_points_in_region_with_hole(tmp_path):
    _geojson(tmp_path / "a.geojson", [_box(0, 0, 100, 100, "+"), _box(40, 40, 60, 60, "-")])
    region = get_geojson_region(str(tmp_path / "a.geojson"))
    assert list(points_in_region(region, [10, 50, 150], [10, 50, 10])) == [True, False, False]   # (col, row)


def test_tracing_path_strips_ome():
    cfg = {"csa_folder": "CSA", "csa_suffix_geojson": "_CSA.geojson"}
    assert tracing_path("/n", "img_001.ome", cfg).name == "img_001_CSA.geojson"
    assert tracing_path("/n", "img_001", cfg).name == "img_001_CSA.geojson"


def test_split_keeps_only_axons_inside(tmp_path, temp_config, log):
    temp_config(CONFIG)
    # Tracing covers columns 0-60: the two left fibers (column 35) are inside.
    nerve, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")])
    assert _counts(info) == (4, 2)
    assert len(kept) == 2 and len(outside) == 2
    assert (kept["y"] < 60).all() and (outside["y"] > 60).all()        # y = column


def test_workbook_sheets_and_summary(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve, *_ = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")])
    wb = openpyxl.load_workbook(nerve / "Morphometrics" / "img_1_morphometrics.xlsx")
    assert wb.sheetnames == ["Summary", "Axon", "Myelin+G-ratio", "Fiber", "Pixel Measurements", "Outside tracing"]
    assert wb["Axon"].max_row == 1 + 2                      # header + 2 axons inside
    assert wb["Pixel Measurements"].max_row == 2 + 2        # note + header + 2
    assert wb["Outside tracing"].max_row == 1 + 2
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(max_row=30) if r[0].value}
    assert str(summary["CSA tracing applied"]).startswith("yes")
    assert (summary["Accepted axons (all)"], summary["Accepted axons outside tracing"]) == (4, 2)


def test_nothing_outside_gives_no_extra_sheet(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 500, 500, "+")])
    assert _counts(info) == (4, 0)
    wb = openpyxl.load_workbook(nerve / "Morphometrics" / "img_1_morphometrics.xlsx")
    assert "Outside tracing" not in wb.sheetnames


def test_exclusion_region_removes_axons_inside_it(tmp_path, temp_config, log):
    temp_config(CONFIG)
    # '+' everywhere, '-' over the top-right fiber (row 35, column 85).
    _, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 500, 500, "+"), _box(70, 20, 100, 50, "-")])
    assert _counts(info) == (4, 1)
    assert (outside["x"].iloc[0], outside["y"].iloc[0]) == pytest.approx((35, 85), abs=1)   # x = row, y = column


def test_positions_follow_resize_scale(tmp_path, temp_config, log):
    temp_config(CONFIG_DOWNSAMPLED)
    # The tracing is in original pixels (240 wide); the segmentation is 120 wide
    # (scale 0.5). Columns 0-120 of the original are columns 0-60 of the segmentation.
    _, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 120, 240, "+")], scale=0.5)
    assert _counts(info) == (4, 2)


def test_no_usable_tracing_changes_nothing(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 60, 120, "-")])     # no '+' region
    assert info is None and len(kept) == 4 and outside is None


def test_nerve_data_counts_inside_and_reports_outside(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve, *_ = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")])
    rows, agg, _ = get_nerve_data(nerve, "40X", log)
    assert rows[0]["total_axons"] == 2 and rows[0]["axons_outside_tracing"] == 2
    assert agg["total_axons"] == 2 and agg["axons_outside_tracing"] == 2
    # CSA = 60 x 120 px at 0.2 µm/px; density uses the inside count.
    assert rows[0]["csa_um2"] == pytest.approx(60 * 120 * 0.04)
    assert rows[0]["axon_density_per_um2"] == pytest.approx(2 / (60 * 120 * 0.04))
    bins = bin_nerve_diameters(nerve / "Morphometrics", "N1", "40X", log)
    assert list(bins["per_image_df"]["Axon Count"]) == [2]                 # distributions agree


def test_all_axons_outside_is_a_measured_zero(tmp_path, temp_config, log):
    temp_config(CONFIG)
    nerve, kept, outside, info = _measure(tmp_path, log, [_box(100, 100, 120, 120, "+")])
    assert _counts(info) == (4, 4) and kept.empty
    rows, agg, _ = get_nerve_data(nerve, "40X", log)
    assert rows[0]["total_axons"] == 0 and rows[0]["axons_outside_tracing"] == 4
    assert rows[0]["csa_um2"] == pytest.approx(20 * 20 * 0.04)
    assert agg["total_axons"] == 0


def test_late_tracing_is_flagged_not_silent(tmp_path, temp_config, log, capsys):
    temp_config(CONFIG)
    # Morphometrics ran without the tracing; the tracing was drawn afterwards.
    nerve, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")], apply=False)
    assert info is None
    rows, agg, _ = get_nerve_data(nerve, "40X", log)
    assert rows[0]["total_axons"] == 4 and rows[0]["axons_outside_tracing"] is None
    assert agg["axons_outside_tracing"] is None
    assert "did not apply" in capsys.readouterr().out


def test_outside_axons_carry_qupath_coordinates(tmp_path, temp_config, log):
    temp_config(CONFIG)
    _, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")])
    # The two right-hand fibers (column 85) are outside; QuPath x is the column.
    assert sorted(outside["qupath_x"]) == pytest.approx([85.5, 85.5], abs=1)
    assert sorted(outside["qupath_y"]) == pytest.approx([35.5, 85.5], abs=1)


def test_qupath_coordinates_follow_resize_scale(tmp_path, temp_config, log):
    temp_config(CONFIG_DOWNSAMPLED)
    _, kept, outside, info = _measure(tmp_path, log, [_box(0, 0, 120, 240, "+")], scale=0.5)
    # Segmented column 85 is column 171 of the original image (2 x 85.5).
    assert sorted(outside["qupath_x"]) == pytest.approx([171.0, 171.0], abs=2)


def test_pooled_file_lists_every_excluded_axon(tmp_path, temp_config, log):
    from morphometrics.tracing import save_outside_tracing
    temp_config(CONFIG)
    n1, _, out1, _ = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")], name="img_1")
    _, _, out2, _ = _measure(tmp_path, log, [_box(0, 0, 60, 120, "+")], name="img_2")
    path = save_outside_tracing([out1, out2, None], n1 / "Morphometrics", "N1")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Outside tracing", "About"]
    header = [c.value for c in wb["Outside tracing"][1]]
    assert header[:7] == ["image", "label", "excluded_because", "qupath_x", "qupath_y", "x", "y"]
    assert wb["Outside tracing"].max_row == 1 + 4                     # 2 images x 2 axons
    assert path.name == "N1_outside_tracing.xlsx"
    assert save_outside_tracing([None, out1.iloc[0:0]], n1 / "Morphometrics", "N1") is None   # nothing outside: no file


def test_qc_figure_draws_tracing_and_excluded_axons(tmp_path, temp_config, log):
    from morphometrics.morphometrics import save_watershed_qc
    temp_config(CONFIG)
    nerve = tmp_path / "N1"
    seg = tmp_path / "img_1_segmented.tif"
    tifffile.imwrite(seg, draw_fibers((120, 120), CENTRES, 15, 9))
    cal = Calibration(0.2, 120, 120, 1.0, "provenance")
    df, raw_df, axon_lbl, fiber_lbl, filtered = get_morphometrics(
        str(seg), "40X", log, return_labels=True, calibration=cal)
    (nerve / "CSA").mkdir(parents=True)
    gj = nerve / "CSA" / "img_1_CSA.geojson"
    _geojson(gj, [_box(0, 0, 60, 120, "+"), _box(10, 10, 20, 20, "-")])
    kept, _, outside, info = split_by_tracing(df, raw_df, gj, 1.0, log)
    save_watershed_qc([seg], [(df, axon_lbl, fiber_lbl, filtered)], nerve / "Morphometrics",
                      "N1", "", "40X", log, tracing_overlays={seg: (info["region_px"], outside)})
    pngs = sorted((nerve / "QC").glob("*.png"))
    assert [p.name for p in pngs] == ["N1_combined_qc.png", "N1_morphometrics_qc.png"]
    assert all(p.stat().st_size > 0 for p in pngs)